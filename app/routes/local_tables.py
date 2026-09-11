from __future__ import annotations

from io import BytesIO
from urllib.parse import quote

from fastapi import APIRouter, Depends, Request, status
from fastapi.responses import PlainTextResponse, RedirectResponse, StreamingResponse
from openpyxl import Workbook
from sqlalchemy import text
from sqlalchemy.exc import SQLAlchemyError
from sqlalchemy.orm import Session

from app.config import get_settings
from app.database import get_db, local_engine
from app.deletion_dependencies import delete_widgets, local_table_dependencies
from app.models import AdminActionLog, Report, ReportExecution
from app.reporting import IDENTIFIER_RE, cleanup_expired_temp_report_tables, quote_identifier, run_select
from app.routes.common import form_lists, render
from app.security import require_importacao, verify_csrf


router = APIRouter()
settings = get_settings()
PREVIEW_LIMIT = 100
BULK_TABLE_ACTIONS = {"export", "clear", "recreate", "delete"}
DESTRUCTIVE_BULK_TABLE_ACTIONS = {"clear", "recreate", "delete"}
INTERNAL_TABLES = {
    "admin_action_logs",
    "auth_logs",
    "dashboard_sources",
    "dashboard_widgets",
    "connector_configs",
    "glpi_import_logs",
    "connector_runs",
    "connector_sync_tables",
    "ldap_configs",
    "portal_groups",
    "portal_group_report_categories",
    "report_categories",
    "report_executions",
    "reports",
    "user_report_categories",
    "user_portal_groups",
    "users",
}


def is_safe_local_user_table(table_name: str) -> bool:
    if not IDENTIFIER_RE.match(table_name):
        return False
    if table_name in INTERNAL_TABLES:
        return False
    if table_name.startswith("glpi_"):
        return False
    return True


def is_safe_report_result_table(table_name: str, db: Session) -> bool:
    if not is_safe_local_user_table(table_name):
        return False
    return db.query(Report).filter(Report.destination_table == table_name).first() is not None


def log_admin_action(
    db: Session,
    user,
    action: str,
    table_name: str | None,
    status_value: str,
    message: str | None,
) -> None:
    db.add(
        AdminActionLog(
            user_id=user.id if user else None,
            username=user.username if user else None,
            action=action,
            table_name=table_name,
            status=status_value,
            message=message,
        )
    )
    db.commit()


def table_metadata_rows(db: Session) -> list[dict]:
    reports_by_table = {
        report.destination_table: report
        for report in db.query(Report).filter(Report.destination_table.isnot(None)).all()
    }
    last_executions = {}
    for execution in (
        db.query(ReportExecution)
        .filter(ReportExecution.destination_table.isnot(None))
        .order_by(ReportExecution.executed_at.desc())
        .all()
    ):
        if execution.destination_table not in last_executions:
            last_executions[execution.destination_table] = execution

    with local_engine.connect() as connection:
        rows = connection.execute(
            text(
                "SELECT TABLE_NAME, TABLE_ROWS, DATA_LENGTH, INDEX_LENGTH, UPDATE_TIME, CREATE_TIME "
                "FROM information_schema.TABLES "
                "WHERE TABLE_SCHEMA = :database_name AND TABLE_TYPE = 'BASE TABLE' "
                "ORDER BY TABLE_NAME"
            ),
            {"database_name": settings.local_db_name},
        ).mappings()
        items = []
        for row in rows:
            table_name = row["TABLE_NAME"]
            report = reports_by_table.get(table_name)
            last_execution = last_executions.get(table_name)
            is_internal = table_name in INTERNAL_TABLES
            is_glpi_copy = table_name.startswith("glpi_")
            is_linked = report is not None
            can_modify = is_safe_report_result_table(table_name, db)
            can_delete = is_safe_local_user_table(table_name)
            if is_internal:
                status_value = "Sistema"
            elif is_glpi_copy:
                status_value = "Copia GLPI"
            elif is_linked:
                status_value = "Resultado"
            else:
                status_value = "Orfa"
            items.append(
                {
                    "table_name": table_name,
                    "report": report,
                    "estimated_rows": row["TABLE_ROWS"] or 0,
                    "estimated_size": (row["DATA_LENGTH"] or 0) + (row["INDEX_LENGTH"] or 0),
                    "updated_at": row["UPDATE_TIME"] or row["CREATE_TIME"],
                    "last_execution": last_execution,
                    "status": status_value,
                    "is_linked": is_linked,
                    "is_orphan": not is_linked and not is_internal and not is_glpi_copy,
                    "can_modify": can_modify,
                    "can_delete": can_delete,
                }
            )
        return items


def filter_tables(items: list[dict], request: Request) -> list[dict]:
    query = request.query_params.get("q", "").strip().lower()
    status_filter = request.query_params.get("status", "").strip()
    linked_filter = request.query_params.get("linked", "").strip()
    filtered = items
    if query:
        filtered = [item for item in filtered if query in item["table_name"].lower()]
    if status_filter:
        filtered = [item for item in filtered if item["status"] == status_filter]
    if linked_filter == "linked":
        filtered = [item for item in filtered if item["is_linked"]]
    elif linked_filter == "orphan":
        filtered = [item for item in filtered if item["is_orphan"]]
    return filtered


def format_bytes(value: int) -> str:
    size = float(value or 0)
    for unit in ["B", "KB", "MB", "GB"]:
        if size < 1024 or unit == "GB":
            return f"{size:.1f} {unit}" if unit != "B" else f"{int(size)} B"
        size /= 1024
    return f"{size:.1f} GB"


def table_exists(table_name: str) -> bool:
    with local_engine.connect() as connection:
        return bool(
            connection.execute(
                text(
                    "SELECT COUNT(*) FROM information_schema.TABLES "
                    "WHERE TABLE_SCHEMA = :database_name AND TABLE_NAME = :table_name"
                ),
                {"database_name": settings.local_db_name, "table_name": table_name},
            ).scalar_one()
        )


def fetch_preview(table_name: str, limit: int = PREVIEW_LIMIT) -> tuple[list[str], list[dict], str | None]:
    if not IDENTIFIER_RE.match(table_name):
        return [], [], "Nome de tabela invalido."
    if not table_exists(table_name):
        return [], [], "Tabela local nao encontrada."
    safe_table = quote_identifier(table_name)
    safe_database = quote_identifier(settings.local_db_name)
    try:
        with local_engine.connect() as connection:
            result = connection.execute(text(f"SELECT * FROM {safe_database}.{safe_table} LIMIT {limit}"))
            columns = list(result.keys())
            rows = [dict(row._mapping) for row in result]
        return columns, rows, None
    except SQLAlchemyError as exc:
        return [], [], f"Nao foi possivel visualizar a tabela: {exc}"


def redirect_with(action: str, table_name: str, success: bool, message: str) -> RedirectResponse:
    key = "message" if success else "error"
    return RedirectResponse(
        f"/admin/local-tables?{key}={quote(message)}",
        status_code=status.HTTP_303_SEE_OTHER,
    )


def local_tables_redirect(success: bool, message: str) -> RedirectResponse:
    key = "message" if success else "error"
    return RedirectResponse(
        f"/admin/local-tables?{key}={quote(message)}",
        status_code=status.HTTP_303_SEE_OTHER,
    )


def selected_table_names_from_form(parsed: dict[str, list[str]]) -> list[str]:
    table_names = []
    for value in parsed.get("table_names", []):
        table_name = value.strip()
        if table_name and IDENTIFIER_RE.match(table_name):
            table_names.append(table_name)
    return list(dict.fromkeys(table_names))


def unique_sheet_title(workbook: Workbook, base: str) -> str:
    title = (base or "Tabela")[:31]
    existing = {sheet.title for sheet in workbook.worksheets}
    if title not in existing:
        return title
    for index in range(2, 1000):
        suffix = f"_{index}"
        candidate = f"{title[:31 - len(suffix)]}{suffix}"
        if candidate not in existing:
            return candidate
    return title[:28] + "_x"


@router.get("/admin/local-tables")
def local_tables(request: Request, db: Session = Depends(get_db)):
    user = require_importacao(request, db)
    if isinstance(user, RedirectResponse):
        return user
    items = table_metadata_rows(db)
    filtered = filter_tables(items, request)
    for item in filtered:
        item["estimated_size_label"] = format_bytes(item["estimated_size"])
    return render(
        request,
        "local_tables.html",
        {
            "active": "local_tables",
            "tables": filtered,
            "filters": {
                "q": request.query_params.get("q", "").strip(),
                "status": request.query_params.get("status", "").strip(),
                "linked": request.query_params.get("linked", "").strip(),
            },
            "message": request.query_params.get("message"),
            "error": request.query_params.get("error"),
        },
        db,
    )


@router.post("/admin/local-tables/bulk", dependencies=[Depends(verify_csrf)])
async def bulk_local_tables(request: Request, db: Session = Depends(get_db)):
    user = require_importacao(request, db)
    if isinstance(user, RedirectResponse):
        return user
    parsed = await form_lists(request)
    action = (parsed.get("bulk_action", [""])[0] or "").strip()
    confirmation = (parsed.get("bulk_confirmation", [""])[0] or "").strip()
    table_names = selected_table_names_from_form(parsed)

    if action not in BULK_TABLE_ACTIONS:
        return local_tables_redirect(False, "Selecione uma acao em massa valida.")
    if not table_names:
        return local_tables_redirect(False, "Selecione ao menos uma tabela.")
    if action in DESTRUCTIVE_BULK_TABLE_ACTIONS and confirmation != "CONFIRMAR":
        return local_tables_redirect(False, f"Acao bloqueada: confirme explicitamente as {len(table_names)} tabelas selecionadas.")

    if action == "export":
        workbook = Workbook()
        workbook.remove(workbook.active)
        success_count = 0
        errors = []
        for table_name in table_names:
            columns, rows, error = fetch_preview(table_name, limit=100000)
            if error:
                errors.append([table_name, error])
                log_admin_action(db, user, "bulk_export", table_name, "error", error)
                continue
            sheet = workbook.create_sheet(unique_sheet_title(workbook, table_name))
            sheet.append(columns)
            for row in rows:
                sheet.append([row.get(column) for column in columns])
            success_count += 1
            log_admin_action(db, user, "bulk_export", table_name, "success", "Exportacao em massa solicitada.")
        if errors or not success_count:
            sheet = workbook.create_sheet(unique_sheet_title(workbook, "Erros"))
            sheet.append(["Tabela", "Resultado"])
            for row in errors:
                sheet.append(row)
        output = BytesIO()
        workbook.save(output)
        output.seek(0)
        return StreamingResponse(
            output,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": 'attachment; filename="tabelas_locais_selecionadas.xlsx"'},
        )

    success_count = 0
    error_count = 0
    for table_name in table_names:
        try:
            if action == "delete":
                can_run_action = is_safe_local_user_table(table_name)
                blocked_message = "Exclusao bloqueada: tabela interna, glpi_* ou nome invalido."
            else:
                can_run_action = is_safe_report_result_table(table_name, db)
                blocked_message = "Acao bloqueada: tabela interna, glpi_*, protegida ou sem relatorio vinculado."
            if not can_run_action:
                message = blocked_message
                log_admin_action(db, user, f"bulk_{action}", table_name, "blocked", message)
                error_count += 1
                continue
            report = db.query(Report).filter(Report.destination_table == table_name).first()
            if action == "clear":
                with local_engine.begin() as connection:
                    connection.execute(text(f"TRUNCATE TABLE {quote_identifier(settings.local_db_name)}.{quote_identifier(table_name)}"))
                message = "Dados limpos em massa mantendo a estrutura."
                log_admin_action(db, user, "bulk_clear", table_name, "success", message)
                success_count += 1
            elif action == "recreate":
                if not report:
                    message = "Recriacao bloqueada: tabela sem relatorio vinculado."
                    log_admin_action(db, user, "bulk_recreate", table_name, "blocked", message)
                    error_count += 1
                    continue
                with local_engine.begin() as connection:
                    connection.execute(text(f"DROP TABLE IF EXISTS {quote_identifier(settings.local_db_name)}.{quote_identifier(table_name)}"))
                execution, _ = run_select(db, report.sql_query, user.id, report.id, report.destination_table, report.modo_salvamento)
                success = execution.status == "success"
                message = "Tabela recriada em massa." if success else execution.error_message or "Falha ao recriar tabela."
                log_admin_action(db, user, "bulk_recreate", table_name, execution.status, message)
                success_count += 1 if success else 0
                error_count += 0 if success else 1
            elif action == "delete":
                with local_engine.begin() as connection:
                    connection.execute(text(f"DROP TABLE {quote_identifier(settings.local_db_name)}.{quote_identifier(table_name)}"))
                message = "Tabela excluida em massa."
                log_admin_action(db, user, "bulk_delete", table_name, "success", message)
                success_count += 1
        except SQLAlchemyError as exc:
            message = str(exc)
            log_admin_action(db, user, f"bulk_{action}", table_name, "error", message)
            error_count += 1

    message = f"Acao em massa concluida. Sucesso: {success_count}. Erros/bloqueios: {error_count}."
    return local_tables_redirect(error_count == 0, message)


@router.post("/admin/local-tables/cleanup-temp-reports", dependencies=[Depends(verify_csrf)])
def cleanup_temp_reports(request: Request, db: Session = Depends(get_db)):
    user = require_importacao(request, db)
    if isinstance(user, RedirectResponse):
        return user
    try:
        removed = cleanup_expired_temp_report_tables(db)
        message = f"Tabelas temporarias expiradas removidas: {removed}."
        log_admin_action(db, user, "cleanup_temp_reports", None, "success", message)
        return RedirectResponse(
            f"/admin/local-tables?message={quote(message)}",
            status_code=status.HTTP_303_SEE_OTHER,
        )
    except SQLAlchemyError as exc:
        message = f"Falha ao limpar tabelas temporarias expiradas: {exc}"
        log_admin_action(db, user, "cleanup_temp_reports", None, "error", message)
        return RedirectResponse(
            f"/admin/local-tables?error={quote(message)}",
            status_code=status.HTTP_303_SEE_OTHER,
        )


@router.get("/admin/local-tables/{table_name}")
def local_table_detail(table_name: str, request: Request, db: Session = Depends(get_db)):
    user = require_importacao(request, db)
    if isinstance(user, RedirectResponse):
        return user
    columns, rows, error = fetch_preview(table_name)
    report = db.query(Report).filter(Report.destination_table == table_name).first()
    return render(
        request,
        "local_table_detail.html",
        {
            "active": "local_tables",
            "table_name": table_name,
            "report": report,
            "columns": columns,
            "rows": rows,
            "error": error,
            "can_modify": is_safe_report_result_table(table_name, db),
            "can_delete": is_safe_local_user_table(table_name),
        },
        db,
        404 if error else 200,
    )


@router.get("/admin/local-tables/{table_name}/history")
def local_table_history(table_name: str, request: Request, db: Session = Depends(get_db)):
    user = require_importacao(request, db)
    if isinstance(user, RedirectResponse):
        return user
    logs = (
        db.query(AdminActionLog)
        .filter(AdminActionLog.table_name == table_name)
        .order_by(AdminActionLog.created_at.desc())
        .limit(100)
        .all()
    )
    return render(
        request,
        "local_table_history.html",
        {"active": "local_tables", "table_name": table_name, "logs": logs},
        db,
    )


@router.get("/admin/local-tables/{table_name}/export")
def export_local_table(table_name: str, request: Request, db: Session = Depends(get_db)):
    user = require_importacao(request, db)
    if isinstance(user, RedirectResponse):
        return user
    columns, rows, error = fetch_preview(table_name, limit=100000)
    if error:
        log_admin_action(db, user, "export", table_name, "error", error)
        return PlainTextResponse(error, status_code=404)
    log_admin_action(db, user, "export", table_name, "success", "Exportacao Excel solicitada.")
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = table_name[:31]
    sheet.append(columns)
    for row in rows:
        sheet.append([row.get(column) for column in columns])
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{table_name}.xlsx"'},
    )


@router.post("/admin/local-tables/{table_name}/refresh", dependencies=[Depends(verify_csrf)])
def refresh_local_table(table_name: str, request: Request, db: Session = Depends(get_db)):
    user = require_importacao(request, db)
    if isinstance(user, RedirectResponse):
        return user
    report = db.query(Report).filter(Report.destination_table == table_name).first()
    if not report:
        message = "Tabela sem relatorio vinculado para atualizacao."
        log_admin_action(db, user, "refresh", table_name, "error", message)
        return redirect_with("refresh", table_name, False, message)
    execution, _ = run_select(db, report.sql_query, user.id, report.id, report.destination_table, report.modo_salvamento)
    success = execution.status == "success"
    message = "Tabela atualizada pelo relatorio vinculado." if success else execution.error_message or "Falha ao atualizar tabela."
    log_admin_action(db, user, "refresh", table_name, execution.status, message)
    return redirect_with("refresh", table_name, success, message)


@router.post("/admin/local-tables/{table_name}/clear", dependencies=[Depends(verify_csrf)])
def clear_local_table(table_name: str, request: Request, db: Session = Depends(get_db)):
    user = require_importacao(request, db)
    if isinstance(user, RedirectResponse):
        return user
    if not is_safe_report_result_table(table_name, db):
        message = "Limpeza bloqueada: apenas tabelas de resultado de relatorios podem ser limpas."
        log_admin_action(db, user, "clear", table_name, "blocked", message)
        return redirect_with("clear", table_name, False, message)
    try:
        with local_engine.begin() as connection:
            connection.execute(text(f"TRUNCATE TABLE {quote_identifier(settings.local_db_name)}.{quote_identifier(table_name)}"))
        message = "Dados limpos mantendo a estrutura da tabela."
        log_admin_action(db, user, "clear", table_name, "success", message)
        return redirect_with("clear", table_name, True, message)
    except SQLAlchemyError as exc:
        message = f"Erro ao limpar tabela: {exc}"
        log_admin_action(db, user, "clear", table_name, "error", message)
        return redirect_with("clear", table_name, False, message)


@router.post("/admin/local-tables/{table_name}/recreate", dependencies=[Depends(verify_csrf)])
def recreate_local_table(table_name: str, request: Request, db: Session = Depends(get_db)):
    user = require_importacao(request, db)
    if isinstance(user, RedirectResponse):
        return user
    report = db.query(Report).filter(Report.destination_table == table_name).first()
    if not report or not is_safe_report_result_table(table_name, db):
        message = "Recriacao bloqueada: tabela sem relatorio vinculado ou protegida."
        log_admin_action(db, user, "recreate", table_name, "blocked", message)
        return redirect_with("recreate", table_name, False, message)
    try:
        with local_engine.begin() as connection:
            connection.execute(text(f"DROP TABLE IF EXISTS {quote_identifier(settings.local_db_name)}.{quote_identifier(table_name)}"))
        execution, _ = run_select(db, report.sql_query, user.id, report.id, report.destination_table, report.modo_salvamento)
        success = execution.status == "success"
        message = "Tabela recriada a partir do SELECT do relatorio." if success else execution.error_message or "Falha ao recriar tabela."
        log_admin_action(db, user, "recreate", table_name, execution.status, message)
        return redirect_with("recreate", table_name, success, message)
    except SQLAlchemyError as exc:
        message = f"Erro ao recriar tabela: {exc}"
        log_admin_action(db, user, "recreate", table_name, "error", message)
        return redirect_with("recreate", table_name, False, message)


@router.post("/admin/local-tables/{table_name}/delete", dependencies=[Depends(verify_csrf)])
def delete_local_table(table_name: str, request: Request, db: Session = Depends(get_db)):
    user = require_importacao(request, db)
    if isinstance(user, RedirectResponse):
        return user
    if not is_safe_local_user_table(table_name):
        message = "Exclusao bloqueada: tabela interna, glpi_* ou nome invalido."
        log_admin_action(db, user, "delete", table_name, "blocked", message)
        return redirect_with("delete", table_name, False, message)
    try:
        deps = local_table_dependencies(db, table_name)
        cascade = request.query_params.get("cascade") == "1"
        if deps["payload"]["has_dependencies"] and not cascade:
            message = "Exclusao bloqueada: confirme a cascata de relatorios, fontes e widgets vinculados."
            log_admin_action(db, user, "delete", table_name, "blocked", message)
            return redirect_with("delete", table_name, False, message)
        widgets = deps["widgets"]
        sources = deps["sources"]
        reports = deps["reports"]
        delete_widgets(db, widgets)
        for source in sources:
            db.delete(source)
        for report in reports:
            report.destination_table = None
        detail = (
            f"Tabela local excluida: {table_name}. "
            f"Relatorios desvinculados: {len(reports)}. Fontes removidas: {len(sources)}. Widgets removidos: {len(widgets)}."
        )
        if reports:
            detail += " Relatorios: " + "; ".join(report.name for report in reports) + "."
        if sources:
            detail += " Fontes: " + "; ".join(source.name for source in sources) + "."
        if widgets:
            detail += " Widgets: " + "; ".join(widget.title for widget in widgets) + "."
        db.flush()
        with local_engine.begin() as connection:
            connection.execute(text(f"DROP TABLE {quote_identifier(settings.local_db_name)}.{quote_identifier(table_name)}"))
        log_admin_action(db, user, "cascade_delete_local_table" if deps["payload"]["has_dependencies"] else "delete", table_name, "success", detail)
        return redirect_with("delete", table_name, True, "Tabela excluida.")
    except SQLAlchemyError as exc:
        message = f"Erro ao excluir tabela: {exc}"
        log_admin_action(db, user, "delete", table_name, "error", message)
        return redirect_with("delete", table_name, False, message)


@router.get("/admin/local-tables/{table_name}/dependencies")
def local_table_delete_dependencies(table_name: str, request: Request, db: Session = Depends(get_db)):
    user = require_importacao(request, db)
    if isinstance(user, RedirectResponse):
        return {"success": False, "error": "Acesso negado."}
    if not is_safe_local_user_table(table_name):
        return {"success": False, "error": "Tabela interna, glpi_* ou nome invalido."}
    return {"success": True, **local_table_dependencies(db, table_name)["payload"]}
