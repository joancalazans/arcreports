from __future__ import annotations

import json
import logging
import re
import threading
import time
from collections import defaultdict
from datetime import datetime, timedelta
from io import BytesIO
from urllib.parse import quote
from zoneinfo import ZoneInfo

from fastapi import APIRouter, BackgroundTasks, Depends, HTTPException, Request, status
from fastapi.responses import HTMLResponse, JSONResponse, PlainTextResponse, RedirectResponse, StreamingResponse
from openpyxl import Workbook
from sqlalchemy import bindparam, text
from sqlalchemy.exc import SQLAlchemyError
from sqlalchemy.orm import Session

from app.config import get_settings
from app.database import get_db, local_engine
from app.deletion_dependencies import delete_widgets, report_dependencies
from app.glpi_import import IDENTIFIER_RE
from app.models import AdminActionLog, DashboardSource, Report, ReportCategory, ReportExecution, SystemConfig, User
from app.reporting import IDENTIFIER_RE as REPORT_IDENTIFIER_RE
from app.reporting import (
    apply_query_timeout,
    clear_report_destination_table,
    extract_local_table_references,
    normalize_save_mode,
    preview_report,
    quote_identifier,
    refresh_report_background,
    run_select,
    validate_select,
)
from app.routes.common import form_data, form_lists, get_allowed_databases, render
from app.security import active_report_categories, can_access_report, log_category_denial, require_relatorios, require_view_user, user_report_category_names, verify_csrf, verify_csrf_header


router = APIRouter()
settings = get_settings()
logger = logging.getLogger(__name__)


CARD_COLORS = {"blue", "green", "amber", "red", "violet", "slate"}
PERIOD_FILTER_MODES = {"nenhum", "filtro_externo"}
SAVE_MODES = {"substituir", "acrescentar"}
BULK_ACTIONS = {"run", "clear", "export", "activate", "deactivate", "delete"}
DESTRUCTIVE_BULK_ACTIONS = {"clear", "delete"}
SCHEDULE_TIME_RE = re.compile(r"^([01]?\d|2[0-3]):([0-5]\d)$")
REPORT_SCHEDULE_HOURS_KEY = "report_schedule_hours"
REPORT_SCHEDULE_ACTIVE_KEY = "report_schedule_active"
ETL_WAIT_BEFORE_REPORTS_KEY = "etl_wait_before_reports"
DEFAULT_REPORT_SCHEDULE_HOURS = "06:30,13:30,20:30"
DEFAULT_REPORT_SCHEDULE_ACTIVE = "true"
DEFAULT_ETL_WAIT_BEFORE_REPORTS = "30"
SCHEDULER_TZ = ZoneInfo("America/Sao_Paulo")
BUILDER_METADATA_CACHE: dict[str, object] = {"expires_at": 0.0, "tables": []}
BUILDER_CACHE_TTL = 300
BUILDER_PREVIEW_LIMIT = 100
BUILDER_ALLOWED_OPERATORS = {
    "eq",
    "ne",
    "contains",
    "starts",
    "ends",
    "gt",
    "lt",
    "between_dates",
    "in_list",
    "empty",
    "not_empty",
}
BUILDER_AGGREGATES = {"count", "sum", "avg", "max", "min"}
BUILDER_SORT_DIRECTIONS = {"asc", "desc"}
BUILDER_EXCLUDED_TABLES = {
    "admin_action_logs",
    "auth_logs",
    "dashboard_sources",
    "dashboard_widgets",
    "connector_configs",
    "glpi_import_logs",
    "connector_runs",
    "connector_sync_tables",
    "ldap_configs",
    "portal_groups",
    "portal_group_report_categories",
    "report_categories",
    "report_executions",
    "reports",
    "user_report_categories",
    "user_portal_groups",
    "users",
}


def clean_optional_text(value: str | None) -> str | None:
    cleaned = (value or "").strip()
    return cleaned or None


def parse_dashboard_order(value: str | None) -> int:
    try:
        return int(value or 100)
    except ValueError:
        return 100


def parse_query_timeout(value: str | None) -> int:
    try:
        return max(1, min(int(value or 60), 600))
    except ValueError:
        return 60


def parse_filter_timeout(value: str | None) -> int:
    try:
        return max(1, min(int(value or 30), 60))
    except ValueError:
        return 30


def default_report_category(db: Session) -> str | None:
    category = db.query(ReportCategory).filter(
        ReportCategory.name == "Departamento 1", ReportCategory.is_active.is_(True)
    ).first()
    if category:
        return category.name
    category = db.query(ReportCategory).filter(ReportCategory.is_active.is_(True)).order_by(ReportCategory.name.asc()).first()
    return category.name if category else None


def validate_report_category(db: Session, category_name: str | None) -> str | None:
    cleaned = clean_optional_text(category_name)
    if not cleaned:
        return default_report_category(db)
    category = db.query(ReportCategory).filter(ReportCategory.name == cleaned, ReportCategory.is_active.is_(True)).first()
    if not category:
        raise ValueError("Categoria de relatorio invalida ou inativa.")
    return category.name


def allowed_builder_schemas() -> set[str]:
    return {settings.local_db_name}


def get_builder_tables(force_refresh: bool = False) -> list[dict]:
    now = time.time()
    if not force_refresh and BUILDER_METADATA_CACHE["expires_at"] > now:
        return BUILDER_METADATA_CACHE["tables"]

    with local_engine.connect() as connection:
        rows = connection.execute(
            text(
                "SELECT TABLE_SCHEMA, TABLE_NAME, TABLE_COMMENT, TABLE_ROWS "
                "FROM information_schema.TABLES "
                "WHERE TABLE_SCHEMA IN :schemas AND TABLE_TYPE IN ('BASE TABLE', 'VIEW') "
                "AND TABLE_NAME NOT IN :excluded_tables "
                "ORDER BY TABLE_SCHEMA, TABLE_NAME"
            ).bindparams(bindparam("schemas", expanding=True), bindparam("excluded_tables", expanding=True)),
            {"schemas": sorted(allowed_builder_schemas()), "excluded_tables": sorted(BUILDER_EXCLUDED_TABLES)},
        ).mappings()
        tables = [
            {
                "schema": row["TABLE_SCHEMA"],
                "name": row["TABLE_NAME"],
                "label": f"{row['TABLE_SCHEMA']}.{row['TABLE_NAME']}",
                "description": row["TABLE_COMMENT"] or "",
                "estimated_rows": row["TABLE_ROWS"] or 0,
            }
            for row in rows
        ]
    BUILDER_METADATA_CACHE["tables"] = tables
    BUILDER_METADATA_CACHE["expires_at"] = now + BUILDER_CACHE_TTL
    return tables


def parse_source(value: str) -> tuple[str, str]:
    if "." not in value:
        raise ValueError("Selecione uma fonte de dados local.")
    schema, table_name = value.split(".", 1)
    if schema not in allowed_builder_schemas() or not REPORT_IDENTIFIER_RE.match(table_name):
        raise ValueError("Fonte de dados invalida.")
    if not any(item["schema"] == schema and item["name"] == table_name for item in get_builder_tables()):
        raise ValueError("Fonte de dados nao encontrada no banco local.")
    return schema, table_name


def get_builder_columns(schema: str, table_name: str) -> list[dict]:
    with local_engine.connect() as connection:
        rows = connection.execute(
            text(
                "SELECT COLUMN_NAME, DATA_TYPE, COLUMN_COMMENT "
                "FROM information_schema.COLUMNS "
                "WHERE TABLE_SCHEMA = :schema AND TABLE_NAME = :table_name "
                "ORDER BY ORDINAL_POSITION"
            ),
            {"schema": schema, "table_name": table_name},
        ).mappings()
        return [
            {
                "name": row["COLUMN_NAME"],
                "data_type": row["DATA_TYPE"],
                "description": row["COLUMN_COMMENT"] or "",
            }
            for row in rows
        ]


def assert_column(column: str, allowed_columns: set[str]) -> str:
    if column not in allowed_columns:
        raise ValueError(f"Coluna invalida: {column}")
    return column


def parse_csv_values(value: str) -> list[str]:
    return [item.strip() for item in (value or "").split(",") if item.strip()]


def sql_literal(value: str) -> str:
    escaped = (value or "").replace("\\", "\\\\").replace("'", "''")
    return f"'{escaped}'"


def build_visual_report(data: dict[str, str], preview_limit: int | None = None) -> tuple[str, dict]:
    schema, table_name = parse_source(data.get("builder_source", ""))
    columns = get_builder_columns(schema, table_name)
    allowed_columns = {column["name"] for column in columns}
    selected_fields = [field for field in data.get("builder_fields", "").split(",") if field]
    selected_fields = [assert_column(field, allowed_columns) for field in selected_fields]
    group_fields = [field for field in data.get("builder_group_fields", "").split(",") if field]
    group_fields = [assert_column(field, allowed_columns) for field in group_fields]
    aggregate_function = data.get("builder_aggregate", "").strip()
    aggregate_field = data.get("builder_aggregate_field", "").strip()

    select_parts = [f"{quote_identifier(field)}" for field in selected_fields]
    if group_fields:
        select_parts = [f"{quote_identifier(field)}" for field in group_fields]
    if aggregate_function:
        if aggregate_function not in BUILDER_AGGREGATES:
            raise ValueError("Agregacao invalida.")
        if aggregate_function == "count":
            select_parts.append("COUNT(*) AS `contagem`")
        else:
            assert_column(aggregate_field, allowed_columns)
            alias = f"{aggregate_function}_{aggregate_field}"
            select_parts.append(f"{aggregate_function.upper()}({quote_identifier(aggregate_field)}) AS {quote_identifier(alias)}")
    if not select_parts:
        raise ValueError("Selecione ao menos um campo.")

    where_parts = []
    for index in range(1, 6):
        field = data.get(f"filter_field_{index}", "").strip()
        operator = data.get(f"filter_operator_{index}", "").strip()
        value = data.get(f"filter_value_{index}", "").strip()
        value_end = data.get(f"filter_value_end_{index}", "").strip()
        if not field or not operator:
            continue
        assert_column(field, allowed_columns)
        if operator not in BUILDER_ALLOWED_OPERATORS:
            raise ValueError("Filtro invalido.")
        column_sql = quote_identifier(field)
        if operator == "eq":
            where_parts.append(f"{column_sql} = {sql_literal(value)}")
        elif operator == "ne":
            where_parts.append(f"{column_sql} <> {sql_literal(value)}")
        elif operator == "contains":
            where_parts.append(f"{column_sql} LIKE {sql_literal('%' + value + '%')}")
        elif operator == "starts":
            where_parts.append(f"{column_sql} LIKE {sql_literal(value + '%')}")
        elif operator == "ends":
            where_parts.append(f"{column_sql} LIKE {sql_literal('%' + value)}")
        elif operator == "gt":
            where_parts.append(f"{column_sql} > {sql_literal(value)}")
        elif operator == "lt":
            where_parts.append(f"{column_sql} < {sql_literal(value)}")
        elif operator == "between_dates":
            where_parts.append(f"{column_sql} BETWEEN {sql_literal(value)} AND {sql_literal(value_end)}")
        elif operator == "in_list":
            values = parse_csv_values(value)
            if not values:
                continue
            where_parts.append(f"{column_sql} IN ({', '.join(sql_literal(item) for item in values)})")
        elif operator == "empty":
            where_parts.append(f"({column_sql} IS NULL OR {column_sql} = '')")
        elif operator == "not_empty":
            where_parts.append(f"({column_sql} IS NOT NULL AND {column_sql} <> '')")

    sql = f"SELECT {', '.join(select_parts)} FROM {quote_identifier(schema)}.{quote_identifier(table_name)}"
    if where_parts:
        sql += " WHERE " + " AND ".join(where_parts)
    if group_fields:
        sql += " GROUP BY " + ", ".join(quote_identifier(field) for field in group_fields)

    sort_field = data.get("builder_sort_field", "").strip()
    sort_direction = data.get("builder_sort_direction", "asc").strip().lower()
    if sort_field:
        assert_column(sort_field, allowed_columns)
        if sort_direction not in BUILDER_SORT_DIRECTIONS:
            sort_direction = "asc"
        sql += f" ORDER BY {quote_identifier(sort_field)} {sort_direction.upper()}"

    limit_raw = data.get("builder_limit", "500")
    try:
        limit = max(1, min(int(limit_raw), 10000))
    except (TypeError, ValueError):
        limit = 500
    if preview_limit is not None:
        limit = min(limit, preview_limit)
    sql += f" LIMIT {limit}"

    config = {
        "source": f"{schema}.{table_name}",
        "fields": selected_fields,
        "group_fields": group_fields,
        "aggregate": aggregate_function,
        "aggregate_field": aggregate_field,
        "sort_field": sort_field,
        "sort_direction": sort_direction,
        "limit": limit,
        "filters": [
            {
                "field": data.get(f"filter_field_{index}", "").strip(),
                "operator": data.get(f"filter_operator_{index}", "").strip(),
                "value": data.get(f"filter_value_{index}", "").strip(),
                "value_end": data.get(f"filter_value_end_{index}", "").strip(),
            }
            for index in range(1, 6)
            if data.get(f"filter_field_{index}", "").strip()
        ],
    }
    return sql, config


def run_builder_preview(sql_query: str) -> tuple[list[str], list[dict], int, int]:
    started = time.perf_counter()
    with local_engine.connect() as connection:
        apply_query_timeout(connection)
        result = connection.execute(text(sql_query))
        mappings = result.mappings().fetchmany(BUILDER_PREVIEW_LIMIT)
        rows = [dict(row) for row in mappings]
        columns = list(result.keys())
    duration_ms = int((time.perf_counter() - started) * 1000)
    return columns, rows, len(rows), duration_ms


def form_context(db: Session, report: Report | None, error: str | None = None, builder_preview: dict | None = None):
    builder_config = {}
    if report and report.builder_config:
        try:
            builder_config = json.loads(report.builder_config)
        except json.JSONDecodeError:
            builder_config = {}
    return {
        "active": "reports",
        "report": report,
        "error": error,
        "builder_tables": get_builder_tables(),
        "builder_config": builder_config,
        "builder_preview": builder_preview,
        "categories": active_report_categories(db),
    }


def apply_report_form(report: Report, data: dict[str, str], name: str, sql_query: str, destination_table: str) -> None:
    card_color = data.get("card_color", "blue").strip()
    report_type = data.get("report_type", "manual").strip()
    modo_salvamento = data.get("modo_salvamento", "substituir").strip()
    report.name = name
    report.description = clean_optional_text(data.get("description"))
    report.sql_query = sql_query
    report.destination_table = destination_table or None
    report.schedule_auto = bool(destination_table) and data.get("schedule_auto") == "on"
    report.schedule_horarios = data.get("schedule_horarios", "").strip() if destination_table else ""
    report.modo_salvamento = modo_salvamento if modo_salvamento in SAVE_MODES else "substituir"
    report.campo_sql_periodo = clean_optional_text(data.get("campo_sql_periodo"))
    modo_filtro_periodo = data.get("modo_filtro_periodo", "filtro_externo").strip()
    report.modo_filtro_periodo = modo_filtro_periodo if modo_filtro_periodo in PERIOD_FILTER_MODES else "filtro_externo"
    report.clausula_periodo_original = None
    report.query_timeout_seconds = parse_query_timeout(data.get("query_timeout_seconds"))
    report.filter_timeout_seconds = parse_filter_timeout(data.get("filter_timeout_seconds"))
    report.report_type = "builder" if report_type == "builder" else "manual"
    report.builder_config = data.get("builder_config") if report.report_type == "builder" else None
    report.is_primary = data.get("is_primary") == "on"
    report.is_active = data.get("is_active") == "on"
    report.show_on_dashboard = data.get("show_on_dashboard") == "on"
    report.dashboard_order = parse_dashboard_order(data.get("dashboard_order"))
    report.category = clean_optional_text(data.get("category"))
    report.card_color = card_color if card_color in CARD_COLORS else "blue"
    report.icon = clean_optional_text(data.get("icon"))
    report.executive_highlight = data.get("executive_highlight") == "on"


def validate_derived_report_source(db: Session, report: Report, sql_query: str) -> None:
    if report.is_primary:
        return
    references = extract_local_table_references(sql_query)
    if not references:
        raise ValueError(
            "Este relatório está tentando usar como fonte a tabela '[tabela]', que é um relatório derivado. "
            "Relatórios derivados devem usar como fonte uma tabela primária (is_primary = 1)."
        )
    report_sources = (
        db.query(Report)
        .filter(Report.destination_table.in_(references))
        .all()
    )
    by_table = {item.destination_table: item for item in report_sources if item.destination_table}
    for table_name in references:
        source = by_table.get(table_name)
        if source and not source.is_primary:
            raise ValueError(
                "Este relatório está tentando usar como fonte a tabela "
                f"'{table_name}', que é um relatório derivado. Relatórios derivados devem usar como fonte "
                "uma tabela primária (is_primary = 1)."
            )
    if not any(source.is_primary for source in report_sources):
        table_name = next((table for table in references if table in by_table), references[0])
        raise ValueError(
            "Este relatório está tentando usar como fonte a tabela "
            f"'{table_name}', que é um relatório derivado. Relatórios derivados devem usar como fonte "
            "uma tabela primária (is_primary = 1)."
        )


def report_destination_rows(report: Report) -> tuple[tuple[list[str], list[dict]] | None, str | None]:
    if not report.destination_table or not IDENTIFIER_RE.match(report.destination_table):
        return None, "Relatorio sem tabela destino local valida para exportacao."
    safe_database = quote_identifier(settings.local_db_name)
    safe_table = quote_identifier(report.destination_table)
    try:
        with local_engine.connect() as connection:
            exists = connection.execute(
                text(
                    "SELECT COUNT(*) FROM information_schema.TABLES "
                    "WHERE TABLE_SCHEMA = :database_name AND TABLE_NAME = :table_name"
                ),
                {"database_name": settings.local_db_name, "table_name": report.destination_table},
            ).scalar_one()
            if not exists:
                return None, "Tabela destino local ainda nao existe."
            result = connection.execute(text(f"SELECT * FROM {safe_database}.{safe_table}"))
            columns = list(result.keys())
            rows = [dict(row._mapping) for row in result]
            return (columns, rows), None
    except SQLAlchemyError as exc:
        return None, f"Nao foi possivel exportar a tabela local: {exc}"


def log_report_admin_action(
    db: Session,
    user,
    action: str,
    report: Report | None,
    status_value: str,
    message: str | None,
) -> None:
    db.add(
        AdminActionLog(
            user_id=user.id if user else None,
            username=user.username if user else None,
            action=action,
            report_id=report.id if report else None,
            report_name=report.name if report else None,
            table_name=report.destination_table if report else None,
            status=status_value,
            message=message,
        )
    )
    db.commit()


def catalog_redirect(success: bool, message: str) -> RedirectResponse:
    key = "message" if success else "error"
    return RedirectResponse(f"/reports?{key}={quote(message)}", status_code=status.HTTP_303_SEE_OTHER)


def selected_report_ids_from_form(parsed: dict[str, list[str]]) -> list[int]:
    ids = []
    for value in parsed.get("report_ids", []):
        try:
            ids.append(int(value))
        except ValueError:
            continue
    return ids


def unique_sheet_title(workbook: Workbook, base: str) -> str:
    cleaned = "".join(char if char.isalnum() or char in {" ", "-", "_"} else "_" for char in base).strip()
    title = (cleaned or "Relatorio")[:31]
    existing = {sheet.title for sheet in workbook.worksheets}
    if title not in existing:
        return title
    for index in range(2, 1000):
        suffix = f"_{index}"
        candidate = f"{title[:31 - len(suffix)]}{suffix}"
        if candidate not in existing:
            return candidate
    return title[:28] + "_x"


def apply_report_visibility(query, user):
    if user.is_admin:
        return query
    allowed_categories = user_report_category_names(user)
    if not allowed_categories:
        return query.filter(False)
    return query.filter(Report.category.in_(allowed_categories))


def category_options_for_user(db: Session, user) -> list[str]:
    if user.is_admin:
        return [category.name for category in db.query(ReportCategory).order_by(ReportCategory.name.asc()).all()]
    return sorted(user_report_category_names(user))


def get_system_config_value(db: Session, key: str, default: str) -> str:
    config = db.get(SystemConfig, key)
    return config.value if config and config.value is not None else default


def set_system_config_value(db: Session, key: str, value: str) -> None:
    config = db.get(SystemConfig, key)
    if config:
        config.value = value
    else:
        db.add(SystemConfig(key=key, value=value))


def parse_schedule_hours(value: str) -> list[str]:
    hours: list[str] = []
    seen: set[str] = set()
    for item in (value or "").split(","):
        raw = item.strip()
        match = SCHEDULE_TIME_RE.match(raw)
        if not match:
            raise ValueError("Informe horarios validos no formato HH:MM separados por virgula.")
        label = f"{int(match.group(1)):02d}:{int(match.group(2)):02d}"
        if label not in seen:
            hours.append(label)
            seen.add(label)
    if not hours:
        raise ValueError("Informe ao menos um horario de execucao.")
    return sorted(hours)


def next_schedule_execution(schedule_hours: str, schedule_active: bool) -> str | None:
    if not schedule_active:
        return None
    try:
        hours = parse_schedule_hours(schedule_hours)
    except ValueError:
        return None
    now = datetime.now(SCHEDULER_TZ)
    candidates = []
    for item in hours:
        hour, minute = item.split(":")
        candidate = now.replace(hour=int(hour), minute=int(minute), second=0, microsecond=0)
        if candidate <= now:
            candidate += timedelta(days=1)
        candidates.append(candidate)
    if not candidates:
        return None
    return min(candidates).strftime("%d/%m/%Y %H:%M")


def scheduler_catalog_context(db: Session) -> dict:
    schedule_hours = get_system_config_value(db, REPORT_SCHEDULE_HOURS_KEY, DEFAULT_REPORT_SCHEDULE_HOURS)
    schedule_active = get_system_config_value(db, REPORT_SCHEDULE_ACTIVE_KEY, DEFAULT_REPORT_SCHEDULE_ACTIVE) == "true"
    etl_wait_raw = get_system_config_value(db, ETL_WAIT_BEFORE_REPORTS_KEY, DEFAULT_ETL_WAIT_BEFORE_REPORTS)
    try:
        etl_wait_before_reports = max(0, int(etl_wait_raw or 0))
    except (TypeError, ValueError):
        etl_wait_before_reports = int(DEFAULT_ETL_WAIT_BEFORE_REPORTS)
    queue_count = (
        db.query(Report)
        .filter(
            Report.is_active.is_(True),
            Report.destination_table.isnot(None),
            Report.destination_table != "",
        )
        .count()
    )
    last_execution = (
        db.query(ReportExecution)
        .join(User, User.id == ReportExecution.user_id)
        .filter(User.username == "scheduler", ReportExecution.status == "success")
        .order_by(ReportExecution.executed_at.desc())
        .first()
    )
    return {
        "schedule_hours": schedule_hours,
        "schedule_active": schedule_active,
        "next_execution": next_schedule_execution(schedule_hours, schedule_active),
        "queue_count": queue_count,
        "last_execution": last_execution.executed_at.strftime("%d/%m/%Y %H:%M") if last_execution else None,
        "etl_wait_before_reports": etl_wait_before_reports,
    }


@router.get("/reports", response_class=HTMLResponse)
def reports(request: Request, db: Session = Depends(get_db)):
    user = require_view_user(request, db)
    if isinstance(user, RedirectResponse):
        return user
    search = request.query_params.get("q", "").strip()
    category = request.query_params.get("category", "").strip()
    status_filter = request.query_params.get("status", "").strip()
    dashboard_filter = request.query_params.get("dashboard", "").strip()

    query = apply_report_visibility(db.query(Report), user)
    if search:
        query = query.filter(Report.name.like(f"%{search}%"))
    if category:
        query = query.filter(Report.category == category)
    if status_filter == "active":
        query = query.filter(Report.is_active.is_(True))
    elif status_filter == "inactive":
        query = query.filter(Report.is_active.is_(False))
    if dashboard_filter == "yes":
        query = query.filter(Report.show_on_dashboard.is_(True))
    elif dashboard_filter == "no":
        query = query.filter(Report.show_on_dashboard.is_(False))

    if search:
        logger.debug(
            "Reports list search SQL q=%r query=%s",
            search,
            query.statement.compile(compile_kwargs={"literal_binds": False}),
        )

    items = query.order_by(Report.name.asc()).all()
    grouped = defaultdict(list)
    for report in items:
        cat = report.category or "Sem categoria"
        grouped[cat].append(report)
    sorted_groups = dict(
        sorted(
            grouped.items(),
            key=lambda item: (
                item[0] == "Sem categoria",
                item[0].lower(),
            ),
        )
    )
    categories = category_options_for_user(db, user)
    return render(
        request,
        "reports.html",
        {
            "active": "reports",
            "reports": items,
            "grouped_reports": sorted_groups,
            "categories": categories,
            "filters": {
                "q": search,
                "category": category,
                "status": status_filter,
                "dashboard": dashboard_filter,
            },
            "message": request.query_params.get("message"),
            "error": request.query_params.get("error"),
            **scheduler_catalog_context(db),
        },
        db,
    )


@router.post("/reports/schedule-config", dependencies=[Depends(verify_csrf)])
async def update_schedule_config(request: Request, db: Session = Depends(get_db)):
    user = require_relatorios(request, db)
    if isinstance(user, RedirectResponse):
        return user
    data = await form_data(request)
    try:
        schedule_hours = ",".join(parse_schedule_hours(data.get("schedule_hours", "")))
    except ValueError as exc:
        return catalog_redirect(False, str(exc))
    schedule_active = "true" if (data.get("schedule_active", "").strip().lower() in {"true", "1", "on", "yes"}) else "false"
    set_system_config_value(db, REPORT_SCHEDULE_HOURS_KEY, schedule_hours)
    set_system_config_value(db, REPORT_SCHEDULE_ACTIVE_KEY, schedule_active)
    db.commit()
    reload_scheduler = getattr(request.app.state, "reload_report_scheduler", None)
    if reload_scheduler:
        reload_scheduler(db)
    log_report_admin_action(
        db,
        user,
        "schedule_config",
        None,
        "success",
        f"Agendamento atualizado: horarios={schedule_hours}; ativo={schedule_active}.",
    )
    return catalog_redirect(True, "Configuracao do agendamento salva com sucesso.")


@router.post("/reports/schedule-run-now", dependencies=[Depends(verify_csrf)])
async def run_schedule_now(request: Request, db: Session = Depends(get_db)):
    user = require_relatorios(request, db)
    if isinstance(user, RedirectResponse):
        return user
    run_all_scheduled_reports = getattr(request.app.state, "run_all_scheduled_reports", None)
    if not run_all_scheduled_reports:
        return catalog_redirect(False, "Scheduler indisponivel.")
    thread = threading.Thread(target=run_all_scheduled_reports, daemon=True)
    thread.start()
    log_report_admin_action(db, user, "schedule_run_now", None, "success", "Execucao manual da fila do scheduler iniciada.")
    return catalog_redirect(True, "Execucao iniciada.")


@router.get("/reports/new", response_class=HTMLResponse)
def new_report(request: Request, db: Session = Depends(get_db)):
    user = require_relatorios(request, db)
    if isinstance(user, RedirectResponse):
        return user
    return render(request, "report_form.html", form_context(db, None), db)


@router.get("/reports/builder/columns")
def builder_columns(request: Request, db: Session = Depends(get_db)):
    user = require_relatorios(request, db)
    if isinstance(user, RedirectResponse):
        return user
    try:
        schema, table_name = parse_source(request.query_params.get("source", ""))
        table_info = next(
            item for item in get_builder_tables() if item["schema"] == schema and item["name"] == table_name
        )
        return JSONResponse({"table": table_info, "columns": get_builder_columns(schema, table_name)})
    except (ValueError, StopIteration) as exc:
        return JSONResponse({"error": str(exc)}, status_code=400)


@router.post("/reports/builder/preview", dependencies=[Depends(verify_csrf_header)])
async def builder_preview(request: Request, db: Session = Depends(get_db)):
    user = require_relatorios(request, db)
    if isinstance(user, RedirectResponse):
        return user
    data = await form_data(request)
    try:
        sql_query, config = build_visual_report(data, preview_limit=BUILDER_PREVIEW_LIMIT)
        columns, rows, row_count, duration_ms = run_builder_preview(sql_query)
        return JSONResponse(
            {
                "columns": columns,
                "rows": rows,
                "row_count": row_count,
                "duration_ms": duration_ms,
                "sql_query": sql_query if user.is_admin else None,
                "config": config,
            }
        )
    except (ValueError, SQLAlchemyError) as exc:
        return JSONResponse({"error": str(exc)}, status_code=400)


@router.post("/reports", dependencies=[Depends(verify_csrf)])
async def create_report(request: Request, db: Session = Depends(get_db)):
    user = require_relatorios(request, db)
    if isinstance(user, RedirectResponse):
        return user
    data = await form_data(request)
    name = data.get("name", "").strip()
    sql_query = data.get("sql_query", "").strip()
    destination_table = data.get("destination_table", "").strip()
    report_type = data.get("report_type", "manual").strip()
    builder_config = None
    if report_type == "builder":
        try:
            sql_query, builder_config = build_visual_report(data)
        except ValueError as exc:
            return render(request, "report_form.html", form_context(db, None, str(exc)), db, 400)
    if not name or not sql_query:
        return render(
            request,
            "report_form.html",
            form_context(db, None, "Informe nome e SQL do relatorio."),
            db,
            400,
        )
    if destination_table and not IDENTIFIER_RE.match(destination_table):
        return render(
            request,
            "report_form.html",
            form_context(db, None, "Tabela destino invalida. Use apenas letras, numeros e underscore."),
            db,
            400,
        )
    try:
        data["category"] = validate_report_category(db, data.get("category")) or ""
    except ValueError as exc:
        return render(request, "report_form.html", form_context(db, None, str(exc)), db, 400)
    try:
        validate_select(sql_query, get_allowed_databases(db))
    except ValueError as exc:
        return render(
            request,
            "report_form.html",
            form_context(db, None, str(exc)),
            db,
            400,
        )
    report = Report(name=name, sql_query=sql_query)
    if builder_config is not None:
        data["builder_config"] = json.dumps(builder_config, ensure_ascii=False)
    apply_report_form(report, data, name, sql_query, destination_table)
    try:
        validate_derived_report_source(db, report, sql_query)
    except ValueError as exc:
        return render(request, "report_form.html", form_context(db, report, str(exc)), db, 400)
    db.add(report)
    db.commit()
    db.refresh(report)
    reload_schedule = getattr(request.app.state, "reload_report_schedule", None)
    if reload_schedule:
        reload_schedule(report.id)
    return RedirectResponse("/reports", status_code=status.HTTP_303_SEE_OTHER)


@router.get("/reports/{report_id}/edit", response_class=HTMLResponse)
def edit_report_page(report_id: int, request: Request, db: Session = Depends(get_db)):
    user = require_relatorios(request, db)
    if isinstance(user, RedirectResponse):
        return user
    report = db.get(Report, report_id)
    if not report:
        return RedirectResponse("/reports", status_code=status.HTTP_303_SEE_OTHER)
    if not can_access_report(user, report):
        log_category_denial(db, user, report, "denied_report_export")
        return PlainTextResponse("Acesso negado para a categoria deste relatorio.", status_code=403)
    return render(request, "report_form.html", form_context(db, report), db)


@router.post("/reports/{report_id}/edit", dependencies=[Depends(verify_csrf)])
async def update_report(report_id: int, request: Request, db: Session = Depends(get_db)):
    user = require_relatorios(request, db)
    if isinstance(user, RedirectResponse):
        return user
    report = db.get(Report, report_id)
    if not report:
        return RedirectResponse("/reports", status_code=status.HTTP_303_SEE_OTHER)
    if not can_access_report(user, report):
        log_category_denial(db, user, report, "denied_report_detail")
        return RedirectResponse("/reports?error=Acesso%20negado%20para%20a%20categoria%20deste%20relatorio.", status_code=status.HTTP_303_SEE_OTHER)

    data = await form_data(request)
    name = data.get("name", "").strip()
    sql_query = data.get("sql_query", "").strip()
    destination_table = data.get("destination_table", "").strip()
    report_type = data.get("report_type", "manual").strip()
    builder_config = None
    if report_type == "builder":
        try:
            sql_query, builder_config = build_visual_report(data)
        except ValueError as exc:
            apply_report_form(report, data, name, report.sql_query, destination_table)
            return render(request, "report_form.html", form_context(db, report, str(exc)), db, 400)
    if not name or not sql_query:
        apply_report_form(report, data, name, sql_query, destination_table)
        return render(
            request,
            "report_form.html",
            form_context(db, report, "Informe nome e SQL do relatorio."),
            db,
            400,
        )
    if destination_table and not IDENTIFIER_RE.match(destination_table):
        apply_report_form(report, data, name, sql_query, destination_table)
        return render(
            request,
            "report_form.html",
            form_context(db, report, "Tabela destino invalida. Use apenas letras, numeros e underscore."),
            db,
            400,
        )
    try:
        data["category"] = validate_report_category(db, data.get("category")) or ""
    except ValueError as exc:
        apply_report_form(report, data, name, sql_query, destination_table)
        return render(request, "report_form.html", form_context(db, report, str(exc)), db, 400)
    try:
        validate_select(sql_query, get_allowed_databases(db))
    except ValueError as exc:
        apply_report_form(report, data, name, sql_query, destination_table)
        return render(
            request,
            "report_form.html",
            form_context(db, report, str(exc)),
            db,
            400,
        )

    if builder_config is not None:
        data["builder_config"] = json.dumps(builder_config, ensure_ascii=False)
    apply_report_form(report, data, name, sql_query, destination_table)
    try:
        validate_derived_report_source(db, report, sql_query)
    except ValueError as exc:
        return render(request, "report_form.html", form_context(db, report, str(exc)), db, 400)
    db.commit()
    reload_schedule = getattr(request.app.state, "reload_report_schedule", None)
    if reload_schedule:
        reload_schedule(report.id)
    return RedirectResponse(f"/reports/{report.id}", status_code=status.HTTP_303_SEE_OTHER)


@router.post("/reports/bulk", dependencies=[Depends(verify_csrf)])
async def bulk_reports(request: Request, db: Session = Depends(get_db)):
    user = require_relatorios(request, db)
    if isinstance(user, RedirectResponse):
        return user
    parsed = await form_lists(request)
    action = (parsed.get("bulk_action", [""])[0] or "").strip()
    report_ids = selected_report_ids_from_form(parsed)
    confirmation = (parsed.get("bulk_confirmation", [""])[0] or "").strip()

    if action not in BULK_ACTIONS:
        return catalog_redirect(False, "Selecione uma acao em massa valida.")
    if not report_ids:
        return catalog_redirect(False, "Selecione ao menos um relatorio.")
    if action in DESTRUCTIVE_BULK_ACTIONS and confirmation != "CONFIRMAR":
        return catalog_redirect(False, f"Acao bloqueada: confirme explicitamente os {len(report_ids)} relatorios selecionados.")

    reports = db.query(Report).filter(Report.id.in_(report_ids)).order_by(Report.name.asc()).all()
    if not reports:
        return catalog_redirect(False, "Nenhum relatorio selecionado foi encontrado.")

    # Validate the entire selection before reading data or applying any action.
    for report in reports:
        if not can_access_report(user, report):
            if action == "export":
                log_report_admin_action(
                    db, user, "export_denied", report, "error",
                    f"Tentativa de exportar relatório {report.id} sem permissão",
                )
            else:
                log_category_denial(db, user, report, "denied_report_bulk")
            raise HTTPException(
                status_code=403,
                detail="Acesso negado para a categoria deste relatório.",
            )

    if action == "export":
        workbook = Workbook()
        workbook.remove(workbook.active)
        errors = []
        exported_count = 0
        for report in reports:
            exported, error = report_destination_rows(report)
            if error:
                errors.append([report.name, report.destination_table or "-", error])
                log_report_admin_action(db, user, "bulk_export", report, "error", error)
                continue
            columns, rows = exported
            sheet = workbook.create_sheet(unique_sheet_title(workbook, report.name))
            sheet.append(columns)
            for row in rows:
                sheet.append([row.get(column) for column in columns])
            exported_count += 1
            log_report_admin_action(db, user, "bulk_export", report, "success", "Exportacao em massa solicitada.")
        if errors or not exported_count:
            sheet = workbook.create_sheet(unique_sheet_title(workbook, "Erros"))
            sheet.append(["Relatorio", "Tabela destino", "Resultado"])
            for row in errors:
                sheet.append(row)
        output = BytesIO()
        workbook.save(output)
        output.seek(0)
        return StreamingResponse(
            output,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": 'attachment; filename="relatorios_selecionados.xlsx"'},
        )

    success_count = 0
    error_count = 0
    for report in reports:
        try:
            if action == "run":
                if not report.is_active:
                    message = "Relatorio inativo nao pode ser executado."
                    log_report_admin_action(db, user, "bulk_run", report, "blocked", message)
                    error_count += 1
                    continue
                execution, _ = run_select(
                    db,
                    report.sql_query,
                    user.id,
                    report.id,
                    report.destination_table,
                    report.modo_salvamento,
                )
                message = (
                    f"Executado com sucesso ({execution.row_count} linhas)."
                    if execution.status == "success"
                    else execution.error_message or "Falha ao executar relatorio."
                )
                log_report_admin_action(db, user, "bulk_run", report, execution.status, message)
                success_count += 1 if execution.status == "success" else 0
                error_count += 0 if execution.status == "success" else 1
            elif action == "clear":
                clear_report_destination_table(db, report.id, report.destination_table)
                message = "Dados da tabela destino limpos mantendo a estrutura."
                log_report_admin_action(db, user, "bulk_clear", report, "success", message)
                success_count += 1
            elif action == "activate":
                report.is_active = True
                db.commit()
                message = "Relatorio ativado em massa."
                log_report_admin_action(db, user, "bulk_activate", report, "success", message)
                success_count += 1
            elif action == "deactivate":
                report.is_active = False
                db.commit()
                message = "Relatorio desativado em massa."
                log_report_admin_action(db, user, "bulk_deactivate", report, "success", message)
                success_count += 1
            elif action == "delete":
                report_name = report.name
                destination_table = report.destination_table
                db.query(ReportExecution).filter(ReportExecution.report_id == report.id).update(
                    {ReportExecution.report_id: None},
                    synchronize_session=False,
                )
                log_report_admin_action(db, user, "bulk_delete", report, "success", "Relatorio excluido em massa.")
                db.delete(report)
                db.commit()
                success_count += 1
        except (ValueError, SQLAlchemyError) as exc:
            db.rollback()
            message = str(exc)
            log_report_admin_action(db, user, f"bulk_{action}", report, "error", message)
            error_count += 1

    message = f"Acao em massa concluida. Sucesso: {success_count}. Erros/bloqueios: {error_count}."
    return catalog_redirect(error_count == 0, message)


@router.get("/reports/{report_id}/export")
def export_report(report_id: int, request: Request, db: Session = Depends(get_db)):
    user = require_view_user(request, db)
    if isinstance(user, RedirectResponse):
        return user
    report = db.get(Report, report_id)
    if not report:
        return RedirectResponse("/reports", status_code=status.HTTP_303_SEE_OTHER)

    if not can_access_report(user, report):
        log_report_admin_action(
            db, user, "export_denied", report, "error",
            f"Tentativa de exportar relatório {report.id} sem permissão",
        )
        raise HTTPException(
            status_code=403,
            detail="Acesso negado para a categoria deste relatório.",
        )

    exported, error = report_destination_rows(report)
    if error:
        return PlainTextResponse(error, status_code=404)
    columns, rows = exported

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Relatorio"
    sheet.append(columns)
    for row in rows:
        sheet.append([row.get(column) for column in columns])

    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    safe_name = "".join(char if char.isalnum() or char in {"-", "_"} else "_" for char in report.name).strip("_")
    filename = f"{safe_name or 'relatorio'}.xlsx"
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@router.post("/reports/{report_id}/toggle", dependencies=[Depends(verify_csrf)])
def toggle_report(report_id: int, request: Request, db: Session = Depends(get_db)):
    user = require_relatorios(request, db)
    if isinstance(user, RedirectResponse):
        return user
    report = db.get(Report, report_id)
    if report:
        report.is_active = not report.is_active
        db.commit()
    return RedirectResponse("/reports", status_code=status.HTTP_303_SEE_OTHER)


@router.post("/reports/{report_id}/toggle-primary", dependencies=[Depends(verify_csrf_header)])
def toggle_primary_report(report_id: int, request: Request, db: Session = Depends(get_db)):
    user = require_relatorios(request, db)
    if isinstance(user, RedirectResponse):
        return JSONResponse({"success": False, "error": "Acesso negado."}, status_code=403)
    report = db.get(Report, report_id)
    if not report:
        return JSONResponse({"success": False, "error": "Relatorio nao encontrado."}, status_code=404)
    if not can_access_report(user, report):
        log_category_denial(db, user, report, "denied_report_toggle_primary")
        return JSONResponse({"success": False, "error": "Acesso negado para a categoria deste relatorio."}, status_code=403)

    report.is_primary = not bool(report.is_primary)
    new_value = bool(report.is_primary)
    db.commit()
    log_report_admin_action(
        db,
        user,
        "toggle_primary",
        report,
        "success",
        f"is_primary={new_value}",
    )
    return JSONResponse({"success": True, "is_primary": new_value, "report_id": report.id})


@router.post("/reports/{report_id}/duplicate", dependencies=[Depends(verify_csrf_header)])
def duplicate_report(report_id: int, request: Request, db: Session = Depends(get_db)):
    user = require_relatorios(request, db)
    if isinstance(user, RedirectResponse):
        return JSONResponse({"success": False, "error": "Acesso negado."}, status_code=403)
    original = db.get(Report, report_id)
    if not original:
        return JSONResponse({"success": False, "error": "Relatorio nao encontrado."}, status_code=404)
    if not can_access_report(user, original):
        log_category_denial(db, user, original, "denied_report_duplicate")
        return JSONResponse({"success": False, "error": "Acesso negado para a categoria deste relatorio."}, status_code=403)

    now = datetime.utcnow()
    new_report = Report(
        name=f"{original.name} (cópia)",
        description=original.description,
        sql_query=original.sql_query,
        destination_table="",
        category=original.category,
        is_active=False,
        is_primary=False,
        created_at=now,
        updated_at=now,
    )
    db.add(new_report)
    db.commit()
    db.refresh(new_report)
    log_report_admin_action(
        db,
        user,
        "duplicate_report",
        new_report,
        "success",
        f"Relatorio duplicado de {original.name}; novo id={new_report.id}.",
    )
    return JSONResponse(
        {
            "success": True,
            "new_id": new_report.id,
            "new_name": new_report.name,
            "redirect": f"/reports/{new_report.id}/edit",
        }
    )


@router.post("/reports/{report_id}/delete", dependencies=[Depends(verify_csrf)])
def delete_report(report_id: int, request: Request, db: Session = Depends(get_db)):
    user = require_relatorios(request, db)
    if isinstance(user, RedirectResponse):
        return user
    report = db.get(Report, report_id)
    if report:
        cascade = request.query_params.get("cascade") == "1"
        deps = report_dependencies(db, report)
        if deps["payload"]["has_dependencies"] and not cascade:
            return catalog_redirect(False, "Exclusao bloqueada: confirme a cascata das fontes e widgets vinculados.")
        widgets = deps["widgets"]
        sources = deps["sources"]
        detail = (
            f"Relatorio excluido: {report.name}. Tabela: {report.destination_table or '-'}. "
            f"Fontes removidas: {len(sources)}. Widgets removidos: {len(widgets)}."
        )
        if sources:
            detail += " Fontes: " + "; ".join(source.name for source in sources) + "."
        if widgets:
            detail += " Widgets: " + "; ".join(widget.title for widget in widgets) + "."
        delete_widgets(db, widgets)
        for source in sources:
            db.delete(source)
        db.query(ReportExecution).filter(ReportExecution.report_id == report.id).update(
            {ReportExecution.report_id: None},
            synchronize_session=False,
        )
        log_report_admin_action(
            db,
            user,
            "report_delete",
            report,
            "success",
            detail,
        )
        db.delete(report)
        db.commit()
    return RedirectResponse("/reports", status_code=status.HTTP_303_SEE_OTHER)


@router.get("/reports/{report_id}/dependencies")
def report_delete_dependencies(report_id: int, request: Request, db: Session = Depends(get_db)):
    user = require_relatorios(request, db)
    if isinstance(user, RedirectResponse):
        return JSONResponse({"success": False, "error": "Acesso negado."}, status_code=403)
    report = db.get(Report, report_id)
    if report is None:
        return JSONResponse({"success": False, "error": "Relatorio nao encontrado."}, status_code=404)
    if not can_access_report(user, report):
        log_category_denial(db, user, report, "denied_report_delete_dependencies")
        return JSONResponse({"success": False, "error": "Acesso negado para a categoria deste relatorio."}, status_code=403)
    return JSONResponse({"success": True, **report_dependencies(db, report)["payload"]})


@router.get("/reports/executions/{execution_id}/status")
def report_execution_status(execution_id: int, request: Request, db: Session = Depends(get_db)):
    user = require_relatorios(request, db)
    if isinstance(user, RedirectResponse):
        return JSONResponse({"error": "Nao autorizado."}, status_code=401)
    execution = db.get(ReportExecution, execution_id)
    if not execution:
        return JSONResponse({"error": "Execucao nao encontrada."}, status_code=404)
    report = db.get(Report, execution.report_id) if execution.report_id else None
    if not can_access_report(user, report):
        log_category_denial(db, user, report, "denied_report_execution_status")
        raise HTTPException(
            status_code=403,
            detail="Acesso negado para a categoria deste relatório.",
        )
    return JSONResponse(
        {
            "execution_id": execution.id,
            "status": execution.status,
            "row_count": execution.row_count,
            "duration_ms": execution.duration_ms,
            "error_message": execution.error_message,
            "finished_at": execution.finished_at.isoformat() if execution.finished_at else None,
        }
    )


@router.get("/reports/{report_id}", response_class=HTMLResponse)
def report_detail(report_id: int, request: Request, db: Session = Depends(get_db)):
    user = require_view_user(request, db)
    if isinstance(user, RedirectResponse):
        return user
    report = db.get(Report, report_id)
    if not report:
        return RedirectResponse("/reports", status_code=status.HTTP_303_SEE_OTHER)
    if not can_access_report(user, report):
        log_category_denial(db, user, report, "denied_report_detail")
        raise HTTPException(
            status_code=403,
            detail="Acesso negado para a categoria deste relatório.",
        )
    executions = (
        db.query(ReportExecution)
        .filter(ReportExecution.report_id == report.id)
        .order_by(ReportExecution.executed_at.desc())
        .limit(10)
        .all()
    )
    return render(
        request,
        "report_detail.html",
        {"active": "reports", "report": report, "executions": executions, "execution": None, "rows": []},
        db,
    )


@router.post(
    "/reports/{report_id}/preview",
    response_class=HTMLResponse,
    dependencies=[Depends(verify_csrf)],
)
def preview_report_endpoint(report_id: int, request: Request, db: Session = Depends(get_db)):
    user = require_relatorios(request, db)
    if isinstance(user, RedirectResponse):
        return user
    report = db.get(Report, report_id)
    if not report or not report.is_active:
        return RedirectResponse(f"/reports/{report_id}", status_code=status.HTTP_303_SEE_OTHER)
    if not can_access_report(user, report):
        log_category_denial(db, user, report, "denied_preview_report_endpoint")
        raise HTTPException(
            status_code=403,
            detail="Acesso negado para a categoria deste relatório.",
        )
    execution, rows = preview_report(db, report.sql_query, user.id, report.id)
    executions = (
        db.query(ReportExecution)
        .filter(ReportExecution.report_id == report.id)
        .order_by(ReportExecution.executed_at.desc())
        .limit(10)
        .all()
    )
    return render(
        request,
        "report_detail.html",
        {
            "active": "reports",
            "report": report,
            "executions": executions,
            "execution": execution,
            "rows": rows,
            "is_preview": True,
        },
        db,
    )


@router.post("/reports/{report_id}/refresh", dependencies=[Depends(verify_csrf)])
def refresh_report_endpoint(
    report_id: int,
    request: Request,
    background_tasks: BackgroundTasks,
    db: Session = Depends(get_db),
):
    user = require_relatorios(request, db)
    if isinstance(user, RedirectResponse):
        return JSONResponse({"error": "Nao autorizado."}, status_code=401)
    report = db.get(Report, report_id)
    if not report or not report.is_active:
        return JSONResponse({"error": "Relatorio nao encontrado."}, status_code=404)
    if not can_access_report(user, report):
        log_category_denial(db, user, report, "denied_refresh_report_endpoint")
        raise HTTPException(
            status_code=403,
            detail="Acesso negado para a categoria deste relatório.",
        )
    if not report.destination_table:
        return JSONResponse({"error": "Relatorio sem tabela destino."}, status_code=400)

    running = (
        db.query(ReportExecution)
        .filter(
            ReportExecution.report_id == report_id,
            ReportExecution.status == "running",
        )
        .first()
    )
    if running:
        return JSONResponse(
            {
                "error": "Ja existe uma atualizacao em andamento para este relatorio.",
                "execution_id": running.id,
            },
            status_code=409,
        )

    execution = ReportExecution(
        report_id=report_id,
        user_id=user.id,
        sql_query=report.sql_query,
        destination_table=report.destination_table,
        status="running",
        executed_at=datetime.utcnow(),
    )
    db.add(execution)
    db.commit()
    db.refresh(execution)
    log_report_admin_action(
        db,
        user,
        "refresh_report_start",
        report,
        "running",
        f"Atualizacao iniciada para {report.name}.",
    )
    background_tasks.add_task(refresh_report_background, report_id, user.id, execution.id)
    return JSONResponse({"execution_id": execution.id, "message": "Atualizacao iniciada."})


@router.post("/reports/{report_id}/run", response_class=HTMLResponse, dependencies=[Depends(verify_csrf)])
def run_report(report_id: int, request: Request, db: Session = Depends(get_db)):
    user = require_relatorios(request, db)
    if isinstance(user, RedirectResponse):
        return user
    report = db.get(Report, report_id)
    if not report:
        return RedirectResponse("/reports", status_code=status.HTTP_303_SEE_OTHER)
    if not can_access_report(user, report):
        log_category_denial(db, user, report, "denied_run_report")
        raise HTTPException(
            status_code=403,
            detail="Acesso negado para a categoria deste relatório.",
        )
    if not report.is_active:
        if request.query_params.get("redirect") == "reports":
            return catalog_redirect(False, "Relatorio inativo nao pode ser executado.")
        return RedirectResponse(f"/reports/{report.id}", status_code=status.HTTP_303_SEE_OTHER)
    execution, rows = run_select(
        db,
        report.sql_query,
        user.id,
        report.id,
        report.destination_table,
        report.modo_salvamento,
    )
    message = (
        f"Relatorio executado com sucesso: {report.name} ({execution.row_count} linhas)."
        if execution.status == "success"
        else execution.error_message or "Falha ao executar relatorio."
    )
    log_report_admin_action(db, user, "run_report", report, execution.status, message)
    if request.query_params.get("redirect") == "reports":
        return catalog_redirect(execution.status == "success", message)
    executions = (
        db.query(ReportExecution)
        .filter(ReportExecution.report_id == report.id)
        .order_by(ReportExecution.executed_at.desc())
        .limit(10)
        .all()
    )
    return render(
        request,
        "report_detail.html",
        {"active": "reports", "report": report, "executions": executions, "execution": execution, "rows": rows},
        db,
    )
