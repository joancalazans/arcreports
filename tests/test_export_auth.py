import asyncio
from io import BytesIO
from unittest.mock import Mock

import pytest
from fastapi import BackgroundTasks, HTTPException
from openpyxl import load_workbook
from starlette.requests import Request

from app.models import AdminActionLog, Report, ReportCategory, ReportExecution
from app.routes import dashboard, reports
from app.security import create_session_token, settings


@pytest.fixture
def report(db_session):
    item = Report(name="Chamados", sql_query="SELECT 1", category="STI",
                  destination_table="resultado", is_active=True, show_on_dashboard=True)
    db_session.add(item)
    db_session.commit()
    return item


def request_for(user):
    headers = [] if user is None else [
        (b"cookie", f"{settings.cookie_name}={create_session_token(user.id)}".encode())
    ]
    return Request({"type": "http", "method": "GET", "path": "/reports/1/export",
                    "query_string": b"", "headers": headers})


def grant_category(db, user, inherited=False):
    category = ReportCategory(name="STI", is_active=True)
    if inherited:
        user.portal_groups[0].report_categories.append(category)
    else:
        user.report_categories.append(category)
    db.commit()


async def response_bytes(response):
    return b"".join([chunk async for chunk in response.body_iterator])


def test_export_denied_without_permission(db_session, view_user, report, monkeypatch):
    """Usuário sem permissão na categoria não consegue exportar o relatório."""
    read = Mock(side_effect=AssertionError("Não pode ler o destino"))
    monkeypatch.setattr(reports, "report_destination_rows", read)
    with pytest.raises(HTTPException) as exc:
        reports.export_report(report.id, request_for(view_user), db_session)
    assert exc.value.status_code == 403
    assert exc.value.detail == "Acesso negado para a categoria deste relatório."
    read.assert_not_called()
    log = db_session.query(AdminActionLog).one()
    assert (log.action, log.status, log.user_id, log.report_id) == (
        "export_denied", "error", view_user.id, report.id)
    assert log.message == f"Tentativa de exportar relatório {report.id} sem permissão"


@pytest.mark.parametrize("permission", ["direct", "group", "admin"])
def test_export_allowed_with_permission(db_session, view_user, report, monkeypatch, permission):
    """Usuário com permissão pode exportar um Excel com os dados esperados."""
    if permission == "admin":
        view_user.is_admin = True
        db_session.commit()
    else:
        grant_category(db_session, view_user, inherited=permission == "group")
    read = Mock(return_value=((["id", "titulo"], [{"id": 1, "titulo": "Teste"}]), None))
    monkeypatch.setattr(reports, "report_destination_rows", read)
    response = reports.export_report(report.id, request_for(view_user), db_session)
    assert response.status_code == 200
    assert response.media_type == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert "Chamados.xlsx" in response.headers["content-disposition"]
    workbook = load_workbook(BytesIO(asyncio.run(response_bytes(response))))
    assert list(workbook.active.values) == [("id", "titulo"), (1, "Teste")]
    read.assert_called_once_with(report)
    assert db_session.query(AdminActionLog).count() == 0


@pytest.mark.parametrize("case", ["anonymous", "missing", "inactive_category", "no_category"])
def test_export_edge_cases(db_session, view_user, report, monkeypatch, case):
    read = Mock(side_effect=AssertionError("Não pode ler o destino"))
    monkeypatch.setattr(reports, "report_destination_rows", read)
    if case == "inactive_category":
        grant_category(db_session, view_user)
        view_user.report_categories[0].is_active = False
        db_session.commit()
    if case == "no_category":
        report.category = None
        db_session.commit()
    if case in {"inactive_category", "no_category"}:
        with pytest.raises(HTTPException) as exc:
            reports.export_report(report.id, request_for(view_user), db_session)
        assert exc.value.status_code == 403
    else:
        response = reports.export_report(
            report.id + 100 if case == "missing" else report.id,
            request_for(None if case == "anonymous" else view_user), db_session)
        assert response.status_code == 303
        assert response.headers["location"] == ("/login" if case == "anonymous" else "/reports")
    read.assert_not_called()


@pytest.mark.parametrize("allowed", [False, True])
def test_bulk_export_authorizes_entire_selection(db_session, view_user, report, monkeypatch, allowed):
    view_user.portal_relatorios = True
    grant_category(db_session, view_user)
    other = Report(name="Z restrito", sql_query="SELECT 2", category="STI" if allowed else "RH")
    db_session.add(other)
    db_session.commit()

    async def form(request):
        return {"bulk_action": ["export"], "report_ids": [str(report.id), str(other.id)]}

    monkeypatch.setattr(reports, "form_lists", form)
    read = Mock(return_value=((["id"], [{"id": 1}]), None))
    monkeypatch.setattr(reports, "report_destination_rows", read)
    if allowed:
        response = asyncio.run(reports.bulk_reports(request_for(view_user), db_session))
        assert response.status_code == 200
        assert read.call_count == 2
        workbook = load_workbook(BytesIO(asyncio.run(response_bytes(response))))
        assert len(workbook.worksheets) == 2
    else:
        with pytest.raises(HTTPException) as exc:
            asyncio.run(reports.bulk_reports(request_for(view_user), db_session))
        assert exc.value.status_code == 403
        read.assert_not_called()
        assert db_session.query(AdminActionLog).one().action == "export_denied"


@pytest.mark.parametrize("endpoint", ["report_detail", "preview_report_endpoint", "run_report",
                                      "refresh_report_endpoint", "report_execution_status"])
def test_other_report_routes_deny_category(db_session, view_user, report, monkeypatch, endpoint):
    view_user.portal_relatorios = True
    db_session.commit()
    for name in ["render", "preview_report", "run_select", "refresh_report_background"]:
        monkeypatch.setattr(reports, name, Mock(side_effect=AssertionError("Acesso indevido")))
    identifier = report.id
    if endpoint == "report_execution_status":
        execution = ReportExecution(report_id=report.id, user_id=view_user.id, sql_query="SELECT 1", status="success")
        db_session.add(execution)
        db_session.commit()
        identifier = execution.id
    extra = {"background_tasks": BackgroundTasks()} if endpoint == "refresh_report_endpoint" else {}
    with pytest.raises(HTTPException) as exc:
        getattr(reports, endpoint)(identifier, request_for(view_user), db=db_session, **extra)
    assert exc.value.status_code == 403
    if extra:
        assert extra["background_tasks"].tasks == []


@pytest.mark.parametrize("endpoint", ["dashboard_report_preview", "search_dashboard_report_detail",
    "dashboard_report_detail", "filter_dashboard_report_detail", "export_dashboard_report_detail",
    "run_dashboard_report"])
def test_dashboard_report_routes_deny_category(db_session, view_user, report, endpoint):
    view_user.portal_relatorios = True
    db_session.commit()
    response = getattr(dashboard, endpoint)(report.id, request_for(view_user), db=db_session)
    assert response.status_code == 403
